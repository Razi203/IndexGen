import pandas as pd
import os
import sys
import glob
import subprocess
import shutil
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Reference, Series

def create_excel_report():
    print("Generating Excel report...")
    
    # Paths
    base_dir = "exps_results"
    cpp_executable = "./extract_and_calc"
    output_xlsx = "clustering_analysis_report.xlsx"
    
    # 1. Find all cluster files
    search_pattern = os.path.join(base_dir, "indices_clusters_*.txt")
    files = glob.glob(search_pattern)
    
    if not files:
        print(f"No files found matching {search_pattern}")
        sys.exit(1)
        
    print(f"Found {len(files)} files to process.")
    
    summary_list = []
    
    # Create Excel Writer
    with pd.ExcelWriter(output_xlsx, engine='openpyxl') as writer:
        
        for file_path in files:
            file_basename = os.path.basename(file_path)
            file_display_name = file_basename.replace("indices_clusters_", "").replace(".txt", "")[:20] # Excel sheet name limit is 31. Suffixes add up to ~9 chars.
            prefix = f"temp_{file_display_name}"
            
            print(f"Processing {file_basename}...")
            
            # 2. Run C++ Extractor
            cmd = [cpp_executable, file_path, prefix]
            result = subprocess.run(cmd, capture_output=True, text=True)
            if result.returncode != 0:
                print(f"Error processing {file_basename}: {result.stderr}")
                continue
                
            # 3. Read CSVs
            clusters_csv = f"{prefix}_clusters.csv"
            hist_csv = f"{prefix}_histogram.csv"
            pairs_csv = f"{prefix}_pairs.csv"
            
            if not all(os.path.exists(f) for f in [clusters_csv, hist_csv, pairs_csv]):
                print(f"Error: Missing output CSVs for {file_basename}")
                continue
                
            df_clusters = pd.read_csv(clusters_csv)
            df_hist = pd.read_csv(hist_csv)
            df_pairs = pd.read_csv(pairs_csv)
            
            # 4. Compute Statistics
            
            # Cluster metrics
            num_clusters = len(df_clusters)
            avg_radius = df_clusters['Radius'].mean()
            avg_size = df_clusters['MemberCount'].mean()
            max_radius = df_clusters['Radius'].max()
            max_size = df_clusters['MemberCount'].max()
            
            # Centroid Distance metrics (from Histogram)
            # Histogram has 'Distance' and 'Count'
            total_pairs_hist = df_hist['Count'].sum()
            if total_pairs_hist > 0:
                min_dist = df_hist[df_hist['Count'] > 0]['Distance'].min()
                max_dist = df_hist[df_hist['Count'] > 0]['Distance'].max()
                weighted_sum = (df_hist['Distance'] * df_hist['Count']).sum()
                avg_dist = weighted_sum / total_pairs_hist
            else:
                min_dist = 0
                max_dist = 0
                avg_dist = 0

            # Percentage Base
            # Total possible pairs N*(N-1)/2
            total_possible_pairs = (num_clusters * (num_clusters - 1)) / 2
            
            # Pairwise metrics
            # df_pairs contains "Close" pairs (dist <= r1+r2+3) which includes Intersections
            
            # Intersection: dist <= r1 + r2
            df_intersections = df_pairs[df_pairs['Distance'] <= (df_pairs['Radius1'] + df_pairs['Radius2'])]
            num_intersections = len(df_intersections)
            
            # Close: dist <= r1 + r2 + 3 (Already filtered in C++)
            num_close = len(df_pairs)
            
            percent_intersections = (num_intersections / total_possible_pairs * 100) if total_possible_pairs > 0 else 0
            percent_close = (num_close / total_possible_pairs * 100) if total_possible_pairs > 0 else 0

            # Append to summary
            summary_list.append({
                'Filename': file_basename,
                'ShortName': file_display_name,
                'NumClusters': num_clusters,
                'AvgRadius': avg_radius,
                'AvgSize': avg_size,
                'MaxRadius': max_radius,
                'MaxSize': max_size,
                'NumIntersections': num_intersections,
                'PercentIntersections': percent_intersections,
                'NumCloseNeighbors': num_close,
                'PercentCloseNeighbors': percent_close,
                'MinDistance': min_dist,
                'MaxDistance': max_dist,
                'AvgDistance': avg_dist
            })
            
            # 5. Write per-file sheets
            # We will create a summary sheet for this file + raw data
            # Sheet: <Name>_Info (Stats + Hist + Pairs?)
            
            # User said "don't add the sheet of clusters".
            # Write to Excel
            # df_clusters.to_excel(writer, sheet_name=f"{file_display_name}_Clusters", index=False)
            
            # Prepare Radius Histogram
            radius_counts = df_clusters['Radius'].value_counts().sort_index()
            df_radius_hist = pd.DataFrame({'Radius': radius_counts.index, 'RadiusCount': radius_counts.values})
            
            # Write Histograms to sheet
            # Distance Hist at Col A (1), Radius Hist at Col D (4)
            sheet_name = f"{file_display_name}_Stats" # Renamed from _Hist to _Stats to reflect content
            
            df_hist.to_excel(writer, sheet_name=sheet_name, startcol=0, index=False)
            df_radius_hist.to_excel(writer, sheet_name=sheet_name, startcol=3, index=False)
            
            ws = writer.sheets[sheet_name]
            
            # Chart 1: Centroid Distance Histogram
            max_row_dist = len(df_hist) + 1
            chart_dist = BarChart()
            chart_dist.title = f"Centroid Distances - {file_display_name}"
            chart_dist.x_axis.title = "Distance"
            chart_dist.y_axis.title = "Count"
            data_dist = Reference(ws, min_col=2, min_row=1, max_row=max_row_dist)
            cats_dist = Reference(ws, min_col=1, min_row=2, max_row=max_row_dist)
            chart_dist.add_data(data_dist, titles_from_data=True)
            chart_dist.set_categories(cats_dist)
            ws.add_chart(chart_dist, "G2")
            
            # Chart 2: Radius Histogram
            max_row_rad = len(df_radius_hist) + 1
            chart_rad = BarChart()
            chart_rad.title = f"Radius Distribution - {file_display_name}"
            chart_rad.x_axis.title = "Radius"
            chart_rad.y_axis.title = "Count"
            data_rad = Reference(ws, min_col=5, min_row=1, max_row=max_row_rad)
            cats_rad = Reference(ws, min_col=4, min_row=2, max_row=max_row_rad)
            
            chart_rad.add_data(data_rad, titles_from_data=True)
            chart_rad.set_categories(cats_rad)
            ws.add_chart(chart_rad, "G20")
            
            # Cleanup temp files
            os.remove(clusters_csv)
            os.remove(hist_csv)
            os.remove(pairs_csv)

        # 6. Global Summary Sheet
        df_summary = pd.DataFrame(summary_list)
        df_summary.to_excel(writer, sheet_name='Comparison', index=False)
        
        # Add Charts to Comparison
        ws_comp = writer.sheets['Comparison']
        
        # Columns mapped to indices (1-based for Excel, df is 0-based)
        # Filename, ShortName, NumClusters, AvgRadius, AvgSize, MaxRadius, MaxSize, NumInter, %Inter, NumClose, %Close, MinDist, MaxDist, AvgDist
        # A, B, C, D, E, F, G, H, I, J, K, L, M, N
        # 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14
        
        # Chart 1: Avg Radius
        chart1 = BarChart()
        chart1.title = "Avg Radius per File"
        chart1.y_axis.title = "Avg Radius"
        data_avg_rad = Reference(ws_comp, min_col=4, min_row=1, max_row=len(df_summary)+1) # Col D
        cats = Reference(ws_comp, min_col=2, min_row=2, max_row=len(df_summary)+1) # ShortName
        chart1.add_data(data_avg_rad, titles_from_data=True)
        chart1.set_categories(cats)
        ws_comp.add_chart(chart1, "P2")
        
        # Chart 2: Avg Distance
        chart2 = BarChart()
        chart2.title = "Avg Centroid Distance per File"
        chart2.y_axis.title = "Avg Distance"
        data_avg_dist = Reference(ws_comp, min_col=14, min_row=1, max_row=len(df_summary)+1) # Col N
        chart2.add_data(data_avg_dist, titles_from_data=True)
        chart2.set_categories(cats)
        ws_comp.add_chart(chart2, "P20")
        
        # Chart 3: Percent Close Neighbors
        chart3 = BarChart()
        chart3.title = "Percent Close Neighbors per File"
        chart3.y_axis.title = "Percentage (%)"
        data_pct_close = Reference(ws_comp, min_col=11, min_row=1, max_row=len(df_summary)+1) # Col K
        chart3.add_data(data_pct_close, titles_from_data=True)
        chart3.set_categories(cats)
        ws_comp.add_chart(chart3, "P38")

    print(f"Report saved to {output_xlsx}")

if __name__ == "__main__":
    create_excel_report()
